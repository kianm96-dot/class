"""교과서 수업 PPT 업그레이드: 원래 장은 그대로 두고, 그 사이에 새 장을 끼워 넣습니다.
사용: python3 tools/upgrade_ppt.py <책 m1|m2> <수업 PPT 폴더> <출력 폴더 | --site>
      --site 이면 수업실 사이트 <책>/<단원>/pptN.pptx 에 바로 씁니다(첫 화면·‘이 차시 자료’의 📥 수업 PPT, 교사만 받음)
      (폴더 안 파일 이름: Ⅰ_1_01_수업 PPT.pptx … 단원(Ⅰ~Ⅳ) 안에서 이름 순서가 차시 순서)
끼워 넣는 장 (모두 원래 PPT의 ‘학습하기’ 틀을 복사해 같은 디자인):
  ✏️ 빈칸 채우기  학습지 요점 정리와 같은 문장 · 같은 번호 빈칸. 클릭마다 답이 하나씩, 이어서 💡 사례
                  → 그 내용을 다룬 원래 ‘학습하기’ 장 바로 뒤 (tools/decks 의 TB 개념 장)
  ➕ 더해 보기    실제 사례 · 딜레마 · 사상가의 말 (수업실 슬라이드의 더해 보기 장) → 관련 빈칸 채우기 장 뒤
  ⏱️ 1분 체크    학습지 1분 체크 OX, 클릭마다 정답 (CHECK)                → 마지막 ‘다음 시간에는’ 장 앞
  🎲 함께 해 보기 그 차시 활동(OX · 분류 · 뒤집기 카드 · 투표 · 수직선 …)  → 1분 체크 뒤
발문 메모는 슬라이드 노트에 들어갑니다. 애니메이션은 ‘클릭할 때 나타내기’라 PowerPoint 슬라이드 쇼에서 동작합니다."""
import sys, re, copy, glob, html, runpy, pathlib, difflib, math
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from lxml import etree

ROOT = pathlib.Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / 'tools'))
from build_deck import compose, acts_of, is_sub, circ

IN = 914400
NS = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main', 'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
      'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
RID = '{%s}' % NS['r']
UNITS = {'Ⅰ': 'self', 'Ⅱ': 'others', 'Ⅲ': 'society', 'Ⅳ': 'nature'}
INK, LAB, RED, BLUE, EXC, DIM = '262626', '843C0C', 'C0392B', '2F6FD0', 'B45F06', '7F7F7F'
X0, Y0, W, H = 0.75, 1.55, 11.85, 5.45          # ‘학습하기’ 틀 안쪽 글 상자 자리(인치)

# ───────── 원래 PPT 읽기
def texts(s): return [sh.text_frame.text.strip() for sh in s.shapes if sh.has_text_frame]
def band(s):
    """‘학습하기’ 장의 초록 띠 제목 상자"""
    c = [sh for sh in s.shapes if sh.has_text_frame and .55 * IN < sh.top < 1.0 * IN and sh.left > 2.4 * IN and sh.width > 5 * IN]
    return c[0] if c else None
def is_learn(s): return '학습하기' in texts(s) and band(s) is not None

def plain(h):
    h = re.sub(r'</(p|h4|li)>|<br\s*/?>', '\n', h or '')
    return re.sub(r'\n{2,}', '\n', html.unescape(re.sub(r'<[^>]+>', '', h))).strip()

# ───────── 글 상자: 문단 = [런...], 런 = dict(t, sz, b, c, step, blank)
def est_h(paras, w):
    """줄 수로 높이 어림(인치): 한글 1em, 그 밖 .55em, 줄 간격 1.3"""
    tot = 0
    for runs, sp in paras:
        em = sum(sum(1 if ord(ch) > 0x2e80 else .55 for ch in r['t']) * r['sz'] for r in runs)
        sz = max(r['sz'] for r in runs)
        lines = max(1, math.ceil(em / (w * 72 * .97)))
        tot += lines * sz * 1.3 / 72 + sp / 72
    return tot

def textbox(slide, paras, x, y, w, h, show=None):
    """show=None: 바탕 상자(단계가 있는 런은 투명, 빈칸은 밑줄만). show=k: k 단계 런만 보이는 덧씌움 상자."""
    tb = slide.shapes.add_textbox(Emu(int(x * IN)), Emu(int(y * IN)), Emu(int(w * IN)), Emu(int(h * IN)))
    tf = tb.text_frame; tf.word_wrap = True
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, 0)
    for i, (runs, sp) in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.3
        if sp: p.space_before = Pt(sp)
        for r in runs:
            rn = p.add_run(); rn.text = r['t']; f = rn.font
            f.size = Pt(r['sz']); f.bold = bool(r.get('b'))
            f.color.rgb = RGBColor.from_string(r.get('c', INK))
            vis = (r.get('step') is None) if show is None else (r.get('step') == show)
            rpr = rn._r.get_or_add_rPr()
            if not vis and r.get('blank') and show is None:
                f.color.rgb = RGBColor.from_string('FFFFFF')   # 바탕 상자의 빈칸: 흰 글자(틀 안이 흰색) + 빨간 밑줄
            elif not vis:
                etree.SubElement(rpr.find('a:solidFill/a:srgbClr', NS), '{%s}alpha' % NS['a'], val='0')
            if r.get('blank') and show is None:
                rpr.set('u', 'sng')
                uf = etree.SubElement(rpr, '{%s}uFill' % NS['a']); sf = etree.SubElement(uf, '{%s}solidFill' % NS['a'])
                etree.SubElement(sf, '{%s}srgbClr' % NS['a'], val=RED)
    return tb

def body(slide, paras, steps):
    """바탕 상자 + 단계마다 덧씌움 상자. 반환: 클릭 순서대로 나타낼 도형 id"""
    textbox(slide, paras, X0, Y0, W, H)
    return [textbox(slide, paras, X0, Y0, W, H, show=k).shape_id for k in range(1, steps + 1)]

def fit(make, limit=H):
    """글자 크기를 줄여 가며 틀 안에 들어가는 문단을 고른다. 들어가지 않으면 None"""
    for sz in (24, 23, 22, 21, 20, 19, 18, 17, 16):
        paras, steps = make(sz)
        if est_h(paras, W) <= limit: return paras, steps
    return None

# ───────── 애니메이션: 클릭마다 하나씩 나타내기(+ 살짝 페이드)
def timing(slide, ids):
    if not ids: return
    n = [2]
    def cid(): n[0] += 1; return n[0]
    clicks = ''
    for sid in ids:
        a, b, c, d, e = cid(), cid(), cid(), cid(), cid()
        clicks += (f'<p:par><p:cTn id="{a}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>'
                   f'<p:par><p:cTn id="{b}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
                   f'<p:par><p:cTn id="{c}" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" grpId="0" nodeType="clickEffect">'
                   f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
                   f'<p:set><p:cBhvr><p:cTn id="{d}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>'
                   f'<p:tgtEl><p:spTgt spid="{sid}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr>'
                   f'<p:to><p:strVal val="visible"/></p:to></p:set>'
                   f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{e}" dur="300"/><p:tgtEl><p:spTgt spid="{sid}"/></p:tgtEl></p:cBhvr></p:animEffect>'
                   f'</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>')
    x = (f'<p:timing xmlns:p="{NS["p"]}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
         f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>{clicks}</p:childTnLst></p:cTn>'
         f'<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
         f'<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
         f'</p:childTnLst></p:cTn></p:par></p:tnLst><p:bldLst>' + ''.join(f'<p:bldP spid="{s}" grpId="0"/>' for s in ids) + '</p:bldLst></p:timing>')
    sld = slide._element
    old = sld.find('p:timing', NS)
    if old is not None: sld.remove(old)
    t = etree.fromstring(x)
    ext = sld.find('p:extLst', NS)
    if ext is not None: ext.addprevious(t)
    else: sld.append(t)

# ───────── 새 장: ‘학습하기’ 틀 복사
def new_slide(prs, base, label, title, pg, note):
    s = prs.slides.add_slide(base.slide_layout)
    for ph in list(s.placeholders): ph._element.getparent().remove(ph._element)
    bg = base._element.find('p:cSld/p:bg', NS)
    if bg is not None: s._element.find('p:cSld', NS).insert(0, copy.deepcopy(bg))
    big = max((sh for sh in base.shapes if sh.has_text_frame and sh.top > 1.3 * IN), key=lambda sh: sh.width * sh.height)
    bd = band(base)
    for sh in base.shapes:
        if sh.shape_id == big.shape_id: continue
        if sh.name.startswith('타원') and sh.width < .4 * IN: continue   # 원래 장의 쪽 표시 점(●○○)은 빼기
        el = copy.deepcopy(sh._element)
        for node in el.iter():
            for k, v in list(node.attrib.items()):
                if k.startswith(RID) and v in base.part.rels:
                    rel = base.part.rels[v]
                    node.set(k, s.part.relate_to(rel.target_ref if rel.is_external else rel.target_part, rel.reltype, rel.is_external))
        s.shapes._spTree.append(el)
        txt = sh.text_frame.text.strip() if sh.has_text_frame else ''
        new = s.shapes[-1]
        if sh.shape_id == bd.shape_id:
            title = re.sub(r'\s*\(\d+(~\d+)?쪽\)', '', title).strip()
            settext(new, title)
            n = sum(1 if ord(ch) > 0x2e80 else .55 for ch in title)
            if n > 19:   # 긴 제목은 띠 안에 들도록 글자를 줄인다
                r0 = new.text_frame.paragraphs[0].runs[0]
                sz = r0.font.size or Pt(28)
                r0.font.size = Pt(max(14, int(sz.pt * 19 / n)))
        elif txt == '학습하기':
            settext(new, label); new.text_frame.word_wrap = False
            from pptx.enum.text import PP_ALIGN
            new.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
        elif re.fullmatch(r'\d+쪽', txt): settext(new, f'{pg}쪽' if pg else '')
    if note: s.notes_slide.notes_text_frame.text = note
    return s

def settext(shape, t):
    p = shape.text_frame.paragraphs[0]
    for extra in shape.text_frame.paragraphs[1:]: extra._p.getparent().remove(extra._p)
    runs = p.runs
    if not runs: p.add_run(); runs = p.runs
    runs[0].text = t
    for r in runs[1:]: r._r.getparent().remove(r._r)

def pg1(pg): return re.match(r'\d+', str(pg)).group(0) if pg else ''

# ───────── 장 만들기
def fill_runs(t, sz, color, ctr, steps, b=False):
    """{{답}} → [번호(보임)] + [빈칸: 투명 글자, 빨간 밑줄 / 덧씌움에서 빨간 답]"""
    out, pos = [], 0
    for m in re.finditer(r'\{\{(.*?)\}\}', t):
        if m.start() > pos: out.append(dict(t=t[pos:m.start()], sz=sz, c=color, b=b))
        ctr[0] += 1; steps[0] += 1
        out.append(dict(t=circ(ctr[0]), sz=sz * .6, c=BLUE, b=True))
        out.append(dict(t=' ' + m.group(1) + ' ', sz=sz, c=RED, b=True, step=steps[0], blank=True))
        pos = m.end()
    if pos < len(t): out.append(dict(t=t[pos:], sz=sz, c=color, b=b))
    return out

def concept_paras(items, ctr0, sz):
    ctr, steps, paras = [ctr0], [0], []
    for it in items:
        lab, txt = it[0], it[1]
        if lab: runs = fill_runs(lab, sz, LAB, ctr, steps, b=True) + [dict(t=' : ', sz=sz, c=LAB, b=True)] + fill_runs(txt, sz, INK, ctr, steps)
        elif is_sub(txt): runs = fill_runs(txt, sz, LAB, ctr, steps, b=True)
        else: runs = fill_runs(txt, sz, INK, ctr, steps)
        paras.append((runs, 7))
        if len(it) > 2:
            steps[0] += 1
            paras.append(([dict(t='💡 ' + it[2], sz=sz - 3, c=EXC, b=True, step=steps[0])], 1))
    return paras, steps[0], ctr[0]

def add_concept(prs, base, x, ctr, note_extra=''):
    """개념 장 하나 → 빈칸 채우기 장(넘치면 둘로 나눔). ctr: 차시 안 빈칸 번호"""
    out = []
    def go(items, h2, note):
        got = fit(lambda sz: concept_paras(items, ctr[0], sz)[:2])
        if got is None and len(items) > 1:
            k = len(items) // 2
            go(items[:k], h2, note); go(items[k:], h2 + ('' if '이어서' in h2 else ' (이어서)'), ''); return
        paras, steps = got or concept_paras(items, ctr[0], 16)[:2]
        ctr[0] = concept_paras(items, ctr[0], 16)[2]
        s = new_slide(prs, base, '빈칸 채우기', h2, pg1(x.get('pg')), note)
        timing(s, body(s, paras, steps)); out.append(s)
    go(x['items'], x['h2'], plain(x.get('note', '')) + note_extra)
    return out

def rows_slide(prs, base, label, title, pg, note, lead, rows, ask=None):
    """lead(늘 보임) + 줄마다 [보이는 글, 클릭하면 나오는 글] + ask. 반환: 장"""
    def make(sz):
        paras, steps = [], 0
        if lead: paras.append(([dict(t=lead, sz=sz, c=LAB, b=True)], 0))
        for i, (main, rev) in enumerate(rows, 1):
            runs = [dict(t=main, sz=sz, c=INK)]
            if rev:
                steps += 1; runs.append(dict(t='   ' + rev, sz=sz - 1, c=RED, b=True, step=steps))
            paras.append((runs, 8))
        if ask:
            steps += 1; paras.append(([dict(t='❓ ' + ask, sz=sz, c=BLUE, b=True, step=steps)], 12))
        return paras, steps
    got = fit(make)
    if got is None:
        if len(rows) > 2:   # 둘로 나눠 싣기
            k = len(rows) // 2
            return rows_slide(prs, base, label, title, pg, note, lead, rows[:k]) + rows_slide(prs, base, label, title + ' (이어서)', pg, '', None, rows[k:], ask)
        got = make(16)
    s = new_slide(prs, base, label, title, pg, note)
    timing(s, body(s, *got))
    return [s]

def extra_slides(prs, base, x, pg):
    """수업실 ➕ 더해 보기 장 → PPT 장 (사진 장은 건너뜀)"""
    t, note = x['type'], plain(x.get('note', ''))
    lad = lambda L: [((a + ' · ' if a else '') + b, None) for a, b in (L or [])]
    if t == 'case':
        return rows_slide(prs, base, '더해 보기', x['h2'].replace('\n', ' '), pg, note, x.get('when'), lad(x.get('facts')), x.get('ask'))
    if t == 'dilemma':
        (an, at), (bn, bt) = x['a'], x['b']
        return rows_slide(prs, base, '더해 보기', '딜레마 · 나라면?', pg, note, x['sit'].replace('\n', ' ') + (f' ({x["small"]})' if x.get('small') else ''),
                          [(f'{an} — {at}', None), (f'{bn} — {bt}', None)], x.get('ask'))
    if t == 'quote':
        return rows_slide(prs, base, '더해 보기', '사상가의 말', pg, note, '“' + x['text'].replace('\n', ' ') + '”', [('— ' + x['who'], None)] + lad(x.get('lad')))
    if t in ('q', 'big'):
        return rows_slide(prs, base, '더해 보기', x['h2'].replace('\n', ' '), pg, note, None, lad(x.get('lad')))
    if t in ('nums', 'trio'):
        return rows_slide(prs, base, '더해 보기', (x.get('h2') or '생각해 보기').replace('\n', ' '), pg, note, None, [(f'{a} · {b}', None) for a, b in x['items']])
    return []

def act_slides(prs, base, acts, blocks, pg):
    """🎲 함께 해 보기: 활동 블록마다 한 장 (모둠 시계는 노트로). 블록이 없으면 활동 이름과 안내 한 장"""
    out = []
    for title, q, _ in acts[:1]:
        bl = [b for b in blocks if b['k'] != 'timer'][:3]
        tm = ''.join(f'\n⏱ {b.get("label", "모둠 활동")} {b["min"]}분' for b in blocks if b['k'] == 'timer')
        if not bl:
            out += rows_slide(prs, base, '함께 해 보기', title, pg, '수업실 활동 화면(act.html)과 같은 활동입니다.' + tm, q,
                              [('① 혼자 생각하기 — 2분 동안 나의 답을 정해 보자', None), ('② 짝·모둠과 나누기 — 서로의 까닭을 들어 보자', None),
                               ('③ 한 문장으로 발표하기 — “나는 ___ 왜냐하면 ___”', None)])
        for b in bl:
            h = b.get('h') or title
            k = b['k']
            if k == 'ox': rows = [(f'{i}. {s}  ( ○ / ✕ )', ('○' if a else '✕') + (f'  {why}' if why else '')) for i, (s, a, why) in enumerate(b['items'], 1)]
            elif k == 'sort': rows = [(f'{i}. {c[0]}', f'→ {b["bins"][c[1]]}' + (f' ({c[2]})' if len(c) > 2 and c[2] else '')) for i, c in enumerate(b['cards'], 1)]
            elif k == 'flip': rows = [(f'🃏 {a}', f'→ {bk}') for a, bk in b['cards']]
            elif k == 'pick': rows = [(f'{chr(65 + i)}. {o[0]}', o[1] if len(o) > 1 else None) for i, o in enumerate(b['opts'])]
            elif k == 'write': rows = [(f'✍️ {a}', f'예) {e}' if e else None) for a, e in b['prompts']]
            elif k == 'vote': rows = [(f'□ {o}   ✋ ____ 명', None) for o in b['opts']]
            elif k == 'line': rows = [(f'“{s}”', None) for s in b['items']] + [(f'{b["left"]}  ① ── ② ── ③ ── ④ ── ⑤  {b["right"]}   (교실 앞 줄에 서 보기)', None)]
            elif k == 'rank': rows = [(f'(   ) {s}', None) for s in b['items']]
            elif k == 'spin': rows = [(f'{i}. {s}', None) for i, s in enumerate(b['items'], 1)]
            else: continue
            lead = {'ox': '맞으면 ○, 틀리면 ✕ — 손으로 먼저 표시해 보자', 'sort': '어느 칸일까? ' + ' · '.join(b.get('bins', [])),
                    'flip': '카드를 뒤집으면?', 'pick': b.get('sit'), 'write': '빈칸을 채워 문장을 완성해 보자', 'vote': '손을 들어 투표해 보자',
                    'line': '나의 생각은 어디쯤? 동의하는 정도에 서 보자', 'rank': '중요한 순서대로 번호를 매겨 보자', 'spin': '모둠이 번호를 하나 골라 미션 수행!'}[k]
            out += rows_slide(prs, base, '함께 해 보기', h, pg, f'{q}{tm}\n수업실 활동 화면(act.html)과 같은 활동입니다.', lead, rows)
    return out

# ───────── 한 차시
def upgrade(path, mod, n, acts, out):
    prs = Presentation(path)
    S = list(prs.slides)
    learn = [(i, band(s).text_frame.text.strip()) for i, s in enumerate(S) if is_learn(s)]
    assert learn, '학습하기 장 없음: ' + str(path)
    base = S[learn[0][0]]
    L = mod['LESSONS'][n - 1]; L = dict(L, slides=compose(L))
    goal = next((x.get('pg') for x in L['tb'] if x['type'] == 'goal'), '')
    # 개념 장과 그 뒤에 이어지는 더해 보기 장
    groups, pre = [], []
    for x in L['slides']:
        if x['type'] == 'concept': groups.append([x, []])
        elif x.get('extra') and x['type'] in ('case', 'dilemma', 'quote', 'q', 'big', 'nums', 'trio'):
            (groups[-1][1] if groups else pre).append(x)
    # 개념 장마다 맞는 원래 ‘학습하기’ 장(제목 닮음, 순서 유지) → 그 장과 같은 제목으로 이어지는 장 뒤
    after, last = {}, 0
    for g in groups:
        h = g[0]['h2'].replace(' (이어서)', '')
        sc = [difflib.SequenceMatcher(None, h, t).ratio() for _, t in learn]
        j = max(range(len(learn)), key=lambda q: sc[q])
        if sc[j] < .45 or j < last: j = last
        last = j
        i, t = learn[j]
        while i + 1 < len(S) and is_learn(S[i + 1]) and band(S[i + 1]).text_frame.text.strip() == t: i += 1
        after.setdefault(i, []).append(g)
    made = {}; ctr = [0]
    for i in sorted(after):
        for c, ex in after[i]:
            made.setdefault(i, []).extend(add_concept(prs, base, c, ctr))
            for x in ex: made[i].extend(extra_slides(prs, base, x, pg1(c.get('pg'))))
    first = learn[0][0] - 1
    for x in pre: made.setdefault(first, []).extend(extra_slides(prs, base, x, pg1(goal)))
    end = []
    chk = mod.get('CHECK', {}).get(n)
    if chk:
        end += rows_slide(prs, base, '1분 체크', '학습지 1분 체크 · 맞으면 ○, 틀리면 ✕', pg1(goal),
                          '학생들이 손으로 ○(머리 위 동그라미) / ✕(팔 엇갈리기)를 먼저 표시한 뒤 클릭해 정답을 보여 준다. ✕ 문장은 어디를 고치면 맞는지 묻는다.',
                          None, [(f'{k}. {q}  ( ○ / ✕ )', '○' if a else '✕') for k, (q, a) in enumerate(chk, 1)])
    blocks = next((A['blocks'] for A in mod.get('ACTS', []) if A['lesson'] == n), [])
    if acts.get(n): end += act_slides(prs, base, acts[n], blocks, pg1(goal))
    made.setdefault(len(S) - 2, []).extend(end)
    # 순서: 원래 장 사이사이에 새 장
    lst = prs.slides._sldIdLst; ids = list(lst)
    byslide = {id(s._element): e for s, e in zip(prs.slides, ids)}
    order = []
    if -1 in made: order += [byslide[id(s._element)] for s in made[-1]]
    for i, e in enumerate(ids[:len(S)]):
        order.append(e)
        order += [byslide[id(s._element)] for s in made.get(i, [])]
    for e in ids: lst.remove(e)
    for e in order: lst.append(e)
    prs.save(out)
    drop_dangling(out)
    return len(S), len(order) - len(S)

def drop_dangling(path):
    """원래 PPT에 대상이 NULL 인 그림 관계가 있으면 python-pptx 가 저장할 때 관계를 지워 svgBlip 참조만 남는다.
    PowerPoint 가 ‘복구’를 묻지 않도록 그 SVG 참조(<a:ext>)만 지운다(보이는 PNG 그림은 그대로)."""
    import zipfile, shutil, os
    z = zipfile.ZipFile(path); files = {n: z.read(n) for n in z.namelist()}; z.close()
    changed = False
    for n in list(files):
        m = re.fullmatch(r'ppt/slides/(slide\d+)\.xml', n)
        if not m: continue
        rel = files.get(f'ppt/slides/_rels/{m.group(1)}.xml.rels', b'')
        ids = set(re.findall(rb'Id="([^"]+)"', rel))
        x = etree.fromstring(files[n]); hit = False
        for sb in x.iter('{http://schemas.microsoft.com/office/drawing/2016/SVG/main}svgBlip'):
            if sb.get(RID + 'embed', '').encode() not in ids:
                ext = sb.getparent(); ext.getparent().remove(ext); hit = True
        if hit: files[n] = etree.tostring(x, xml_declaration=True, encoding='UTF-8', standalone=True); changed = True
    if changed:
        tmp = str(path) + '.tmp'
        with zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as z:
            for n, d in files.items(): z.writestr(n, d)
        os.replace(tmp, path)

def main(book, src, dst):
    site = dst == '--site'
    if not site: dst = pathlib.Path(dst); dst.mkdir(parents=True, exist_ok=True)
    files = sorted(glob.glob(str(pathlib.Path(src) / '*.pptx')))
    for r, unit in UNITS.items():
        fs = [f for f in files if pathlib.Path(f).name.startswith(r + '_')]
        if not fs: continue
        mod = runpy.run_path(str(ROOT / 'tools/decks' / f'{book}_{unit}.py'))
        acts = acts_of(book, unit)
        for n, f in enumerate(fs, 1):
            out = ROOT / book / unit / f'ppt{n}.pptx' if site else dst / pathlib.Path(f).name.replace('수업 PPT', '수업 PPT 업그레이드')
            a, b = upgrade(f, mod, n, acts, out)
            print(f'{out.name}  원래 {a}장 + 새 {b}장')

if __name__ == '__main__':
    main(*sys.argv[1:4])
